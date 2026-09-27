"""Conversation-scoped artifacts, safe filenames and attachment-only downloads."""

import csv
import io
import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

from .db import now, uid


class Artifacts:
    def __init__(self, root, store):
        self.root = Path(root) / "artifacts"
        self.root.mkdir(exist_ok=True, mode=0o700)
        self.store = store

    def save(self, conversation_id, run_id, name, data, mime):
        name = re.sub(r"[^\w. -]", "_", name, flags=re.UNICODE).strip(". ")[:120] or "file.txt"
        if len(data) > 50_000_000:
            raise ValueError("Artifact exceeds 50 MB")
        artifact_id = uid()
        path = self.root / artifact_id
        path.write_bytes(data)
        self.store.execute(
            "INSERT INTO artifacts VALUES(?,?,?,?,?,?,?,?)",
            (artifact_id, conversation_id, run_id, name, mime, str(path), len(data), now()),
        )
        return {
            "id": artifact_id,
            "name": name,
            "size": len(data),
            "mime": mime,
            "url": f"/api/artifacts/{artifact_id}/download",
        }

    def get(self, artifact_id, conversation_id=None):
        item = self.store.one("SELECT * FROM artifacts WHERE id=?", (artifact_id,))
        if not item or (conversation_id and item["conversation_id"] != conversation_id):
            raise ValueError("Artifact not found in this conversation")
        if not Path(item["path"]).resolve().is_relative_to(self.root.resolve()):
            raise ValueError("Invalid artifact location")
        return item

    def text(self, artifact_id, conversation_id):
        item = self.get(artifact_id, conversation_id)
        if item["size"] > 5_000_000:
            raise ValueError("File is too large for text extraction")
        data = Path(item["path"]).read_bytes()
        suffix = Path(item["name"]).suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n".join(p.extract_text() or "" for p in reader.pages[:20])[:30000]
        if suffix not in (
            ".txt",
            ".md",
            ".json",
            ".csv",
            ".py",
            ".html",
            ".obj",
            ".yaml",
            ".yml",
            ".js",
            ".css",
        ):
            raise ValueError("Read supports UTF-8 text and the first 20 pages of text PDFs")
        return data.decode("utf-8", errors="replace")[:30000]

    def document(self, conversation_id, run_id, name, fmt, content):
        buffer = io.BytesIO()
        name = Path(name).stem + "." + fmt
        if fmt == "docx":
            from docx import Document
            from docx.shared import Pt

            document = Document()
            document.styles["Normal"].font.name = "Calibri"
            document.styles["Normal"].font.size = Pt(11)
            for line in content.splitlines():
                match = re.match(r"^(#{1,3})\s+(.+)$", line)
                if match:
                    document.add_heading(match[2], len(match[1]))
                elif line.startswith("- "):
                    document.add_paragraph(line[2:], style="List Bullet")
                else:
                    document.add_paragraph(line)
            document.save(buffer)
            data, mime = (
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        elif fmt == "pdf":
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            styles = getSampleStyleSheet()
            for style in styles.byName.values():
                style.fontName = "STSong-Light"
                style.wordWrap = "CJK"
            elements = []
            for line in content.splitlines():
                heading = re.match(r"^(#{1,3})\s+(.+)$", line)
                style = styles[f"Heading{len(heading[1])}"] if heading else styles["BodyText"]
                elements.append(
                    Paragraph(escape(heading[2] if heading else line) or "<br/>", style)
                )
                elements.append(Spacer(1, 4))
            SimpleDocTemplate(buffer, title=Path(name).stem, rightMargin=48, leftMargin=48).build(
                elements
            )
            data, mime = buffer.getvalue(), "application/pdf"
        else:
            if fmt == "json":
                content = json.dumps(json.loads(content), ensure_ascii=False, indent=2)
            if fmt == "csv":
                # Prevent spreadsheet formula execution when a user opens generated CSV.
                output = io.StringIO()
                writer = csv.writer(output)
                for row in csv.reader(io.StringIO(content)):
                    writer.writerow(
                        [
                            "'" + c
                            if c.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
                            else c
                            for c in row
                        ]
                    )
                content = output.getvalue()
            data = content.encode("utf-8")
            mime = {
                "txt": "text/plain",
                "md": "text/markdown",
                "json": "application/json",
                "csv": "text/csv",
                "html": "text/html",
            }[fmt]
        return self.save(conversation_id, run_id, name, data, mime)
